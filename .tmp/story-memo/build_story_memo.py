from __future__ import annotations

import json
import math
import re
import shutil
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"D:\TestProj\SDS_Games")
TMP = ROOT / ".tmp" / "story-memo"
DATA_PATH = TMP / "story-data.json"
TEMPLATE = Path(r"D:\Users\s494326\.codex\plugins\cache\openai-curated-remote\openai-templates\0.1.0\skills\artifact-template-legal-memorandum\assets\reference.docx")
OUTPUT = ROOT / "SDS_Games_雙周目劇情因果審查稿.docx"
MAP_DIR = TMP / "maps"
FONT_SERIF = Path(r"C:\Windows\Fonts\NotoSerifTC-VF.ttf")
FONT_SANS = Path(r"C:\Windows\Fonts\NotoSansTC-VF.ttf")

CHAPTER_TITLES = {
    1: "南門以外",
    2: "斷路上的藥",
    3: "影子仍在守望",
    4: "石心與灰雨",
    5: "元素失去形狀",
    6: "龍守著封痕",
    7: "魔王墜落之處",
}

MAIN_BOSSES = {
    1: ("forest_guardian", "森林守衛"),
    2: ("lich", "巫妖"),
    3: ("shadow_commander", "暗影指揮官"),
    4: ("ancient_titan", "遠古泰坦"),
    5: ("elemental_lord", "元素領主"),
    6: ("elder_dragon", "龍族長者"),
    7: ("demon_lord_asariel", "魔王阿薩列爾"),
}

ROUTE_BOSSES = {
    1: [("ambush_mantis", "伏獵螳螂", "ch1_s07_silver_snare")],
    2: [], 3: [], 4: [], 5: [], 6: [], 7: [],
}

EXTERNAL_BOSSES = {
    1: ("prologue_blood_moon_stag", "前導角鹿再戰"),
    2: ("nameless_curse", "無名之咒"),
    3: ("expedition_supreme_commander", "遠征軍最高指揮官"),
    4: ("ash_baron_external", "灰燼男爵"),
    5: ("entropy_balance_external", "熵"),
    6: ("light_trial_external", "光明試煉"),
    7: ("void_revelation_external", "虛空 Boss"),
}

ACTOR_NAMES = {
    "player": "主角",
    "village_elder": "村長",
    "town_scholar": "伊萊",
    "herbalist": "米婭",
    "standard_bearer_frey": "芙蕾",
    "lamplighter_tavi": "塔維",
    "blacksmith": "鐵匠",
    "street_beggar": "艾洛／街頭乞者",
    "young_ailo": "年幼艾洛",
    "neelu": "妮露",
    "casino_owner": "維斯珀",
    "casino_dealer": "賭場荷官",
    "merchant": "商人",
    "black_market": "黑市商人",
    "elder_dragon": "龍族長者",
    "demon_lord_asariel": "魔王阿薩列爾",
    "lich": "巫妖",
    "shadow_commander": "暗影指揮官",
    "drowned_oracle": "溺亡預言者",
    "thorn_witch": "荊棘女巫",
}

CORE_CHARACTER_IDS = [
    "village_elder", "herbalist", "town_scholar", "standard_bearer_frey",
    "lamplighter_tavi", "blacksmith", "street_beggar", "casino_owner"
]

COLORS = {
    "ink": "#252525",
    "muted": "#666666",
    "rule": "#B8B8B8",
    "gold": "#B9892E",
    "gold_fill": "#FFF4D5",
    "blue": "#246A92",
    "blue_fill": "#E8F3FA",
    "red": "#7A2626",
    "red_fill": "#F5E4E4",
    "gray": "#8A8A8A",
    "gray_fill": "#F1F1F1",
    "amber": "#D08A13",
    "paper": "#FAF9F6",
}


def load_data():
    return json.loads(DATA_PATH.read_text(encoding="utf-8"))


def concise(text, limit=360):
    if text is None:
        return "未在資料中明定。"
    text = re.sub(r"\s+", " ", str(text)).strip()
    return text if len(text) <= limit else text[:limit - 1].rstrip() + "…"


def scene_title(scene):
    """Runtime currently names Chapter 1 scenes only; keep later ids visible without inventing canon titles."""
    return scene.get("title") or f"未定標題（{scene['id'].split('_', 2)[-1]}）"


def beat_summary(scene, condition, limit=420):
    beats = [b for b in scene.get("beats", []) if b.get("condition") in ("any", condition)]
    fragments = []
    for beat in beats:
        text = concise(beat.get("text"), 150)
        if not text:
            continue
        if beat.get("actorId"):
            text = f"{ACTOR_NAMES.get(beat['actorId'], beat['actorId'])}：{text}"
        fragments.append(text)
    return concise(" ／ ".join(fragments), limit)


def scene_actor_ids(scene):
    ids = []
    for beat in scene.get("beats", []):
        aid = beat.get("actorId")
        if aid and aid not in ids:
            ids.append(aid)
    return ids


def scene_run_profile(scene):
    conds = {b.get("condition") for b in scene.get("beats", [])}
    return {
        "shared": "any" in conds,
        "first": "first_run" in conds,
        "second": "second_run" in conds,
    }


def audit_scene(scene):
    missing = []
    for key, label in (("entry", "前置觸發"), ("objective", "場景目的"), ("outputsRaw", "輸出"), ("exit", "承接")):
        if not scene.get(key):
            missing.append(label)
    if missing:
        return "需補強", "缺少「" + "、".join(missing) + "」欄位。"
    profile = scene_run_profile(scene)
    if scene["chapter"] >= 5 and not (profile["first"] or profile["second"]):
        return "需補強", "後段場景只有共用節拍；若預期兩周目改變結果，需確認差異是否由其他場景承擔。"
    return "已確立", "runtime 已具備入口、目的、輸出與離場承接。"


def extract_md_scene_ids():
    path = ROOT / "docs" / "MAIN_STORY_BIBLE.md"
    text = path.read_text(encoding="utf-8", errors="replace")
    return sorted(set(re.findall(r"ch[1-7]_s\d{2}_[a-z0-9_]+", text)))


def font(size, bold=False):
    # Variable Noto fonts render Chinese reliably through Pillow.
    return ImageFont.truetype(str(FONT_SANS), size=size)


def wrap_px(draw, text, fnt, max_width, max_lines=None):
    text = concise(text, 500)
    lines, current = [], ""
    for ch in text:
        trial = current + ch
        if draw.textbbox((0, 0), trial, font=fnt)[2] <= max_width or not current:
            current = trial
        else:
            lines.append(current)
            current = ch
            if max_lines and len(lines) >= max_lines:
                break
    if current and (not max_lines or len(lines) < max_lines):
        lines.append(current)
    if max_lines and len(lines) == max_lines and len("".join(lines)) < len(text):
        lines[-1] = lines[-1][:-1] + "…"
    return lines


def draw_arrow(draw, start, end, color, width=7, dotted=False):
    x1, y1 = start
    x2, y2 = end
    if dotted:
        segments = 14
        for i in range(0, segments, 2):
            a, b = i / segments, min((i + 1) / segments, 1)
            draw.line((x1 + (x2-x1)*a, y1 + (y2-y1)*a, x1 + (x2-x1)*b, y1 + (y2-y1)*b), fill=color, width=width)
    else:
        draw.line((x1, y1, x2, y2), fill=color, width=width)
    angle = math.atan2(y2-y1, x2-x1)
    size = 24
    pts = [
        (x2, y2),
        (x2 - size*math.cos(angle-0.55), y2 - size*math.sin(angle-0.55)),
        (x2 - size*math.cos(angle+0.55), y2 - size*math.sin(angle+0.55)),
    ]
    draw.polygon(pts, fill=color)


def draw_box(draw, box, title, subtitle="", kind="main", marker=None):
    x1, y1, x2, y2 = box
    palette = {
        "main": (COLORS["paper"], COLORS["ink"]),
        "side": (COLORS["gold_fill"], COLORS["gold"]),
        "evidence": (COLORS["blue_fill"], COLORS["blue"]),
        "boss": (COLORS["red_fill"], COLORS["red"]),
        "pending": (COLORS["gray_fill"], COLORS["gray"]),
    }
    fill, outline = palette[kind]
    draw.rounded_rectangle(box, radius=20, fill=fill, outline=outline, width=5)
    tf = font(31)
    sf = font(24)
    title_lines = wrap_px(draw, title, tf, x2-x1-42, 2)
    y = y1 + 22
    for line in title_lines:
        draw.text((x1+21, y), line, font=tf, fill=COLORS["ink"])
        y += 40
    if subtitle:
        for line in wrap_px(draw, subtitle, sf, x2-x1-42, 2):
            draw.text((x1+21, y+6), line, font=sf, fill=COLORS["muted"])
            y += 32
    if marker:
        draw.ellipse((x2-48, y1-18, x2-8, y1+22), fill=COLORS["amber"])
        draw.text((x2-36, y1-13), "!", font=font(25), fill="white")


def map_header(draw, title, subtitle, width):
    draw.text((90, 65), title, font=font(58), fill=COLORS["ink"])
    draw.text((90, 140), subtitle, font=font(29), fill=COLORS["muted"])
    draw.line((90, 196, width-90, 196), fill=COLORS["gold"], width=5)


def map_legend(draw, x, y):
    items = [
        ("main", "已確立主線"), ("side", "可選人物／多人支線"),
        ("evidence", "證據／認知改變"), ("boss", "Boss／不可逆代價"),
        ("pending", "待定／暫緩"),
    ]
    for idx, (kind, label) in enumerate(items):
        bx = x + idx * 535
        draw_box(draw, (bx, y, bx+55, y+55), "", kind=kind)
        draw.text((bx+72, y+7), label, font=font(24), fill=COLORS["muted"])
    draw.ellipse((x+5*535, y+8, x+5*535+38, y+46), fill=COLORS["amber"])
    draw.text((x+5*535+52, y+7), "需補強", font=font(24), fill=COLORS["muted"])


def create_overview_map(data):
    w, h = 3600, 2200
    img = Image.new("RGB", (w, h), COLORS["paper"])
    d = ImageDraw.Draw(img)
    map_header(d, "七章雙周目因果總覽", "主幹、章節 Boss、第一周目假勝利與第二周目修正", w)
    x0, gap, bw = 95, 45, 455
    y_first, y_second = 370, 1130
    d.text((95, 270), "第一周目｜資訊不足下仍然成立的勝利", font=font(34), fill=COLORS["red"])
    d.text((95, 1030), "第二周目｜保留記憶理解、重走物理路線", font=font(34), fill=COLORS["blue"])
    first_centers, second_centers = [], []
    for ch in range(1, 8):
        x = x0 + (ch-1)*(bw+gap)
        scenes = data["sceneOrderByChapter"][str(ch)]
        main_id, main_name = MAIN_BOSSES[ch]
        draw_box(d, (x, y_first, x+bw, y_first+300), f"第 {ch} 章｜{CHAPTER_TITLES[ch]}", f"{len(scenes)} 場 · {main_name}", "boss")
        first_centers.append((x+bw, y_first+150))
        ext_id, ext_name = EXTERNAL_BOSSES[ch]
        draw_box(d, (x, y_second, x+bw, y_second+300), f"第 {ch} 章重讀", f"外部 Boss 保留位：{ext_name}", "pending")
        second_centers.append((x+bw, y_second+150))
        draw_arrow(d, (x+bw/2, y_first+300), (x+bw/2, y_second), COLORS["blue"], width=6)
        if ch < 7:
            nx = x0 + ch*(bw+gap)
            draw_arrow(d, (x+bw, y_first+150), (nx, y_first+150), COLORS["ink"], width=7)
            draw_arrow(d, (x+bw, y_second+150), (nx, y_second+150), COLORS["gray"], width=6, dotted=True)
    draw_box(d, (1040, 1700, 2560, 1950), "第一周目結束：魔王戰鬥體倒下，但根、回聲與生命層未被完整處理", "真實擊殺尚未完成；玩家取得重新理解既有證據的入口。", "boss")
    draw_arrow(d, (3370, y_first+300), (2560, 1820), COLORS["red"], width=8)
    draw_box(d, (2630, 1700, 3500, 1950), "第二周目真結局", "以當周目重新取得的生命種子、古代符文、森林精華與微光閉合核心。", "evidence")
    draw_arrow(d, (3370, y_second+300), (3060, 1700), COLORS["blue"], width=8)
    map_legend(d, 95, 2060)
    path = MAP_DIR / "00_overview.png"
    img.save(path, quality=95)
    return path


def create_trunk_map(chapter, scenes):
    w, h = 3600, 2250
    img = Image.new("RGB", (w, h), COLORS["paper"])
    d = ImageDraw.Draw(img)
    map_header(d, f"第 {chapter} 章主幹｜{CHAPTER_TITLES[chapter]}", "上軌：第一周目　下軌：第二周目；同場景的條件節拍分列呈現", w)
    cols = 6
    bw, bh = 500, 230
    xgap, ygap = 55, 85
    x0 = 90
    lanes = [("第一周目", 300, "first_run"), ("第二周目", 1215, "second_run")]
    for lane_name, y0, cond in lanes:
        d.text((x0, y0-70), lane_name, font=font(33), fill=COLORS["red"] if cond == "first_run" else COLORS["blue"])
        centers = []
        for idx, s in enumerate(scenes):
            row, col = divmod(idx, cols)
            x = x0 + col*(bw+xgap)
            y = y0 + row*(bh+ygap)
            status, _ = audit_scene(s)
            is_boss = MAIN_BOSSES[chapter][0] in (s["id"] + " " + scene_title(s)).lower() or "boss" in s["objective"].lower()
            kind = "boss" if is_boss or s["id"].endswith(("forest_guardian", "keeper_of_names", "shadow_commander", "titan_rises", "elemental_lord", "dragon_convergence", "last_core")) else "main"
            profile = scene_run_profile(s)
            if cond == "second_run" and not profile["second"] and chapter >= 5:
                subtitle = "共用節拍；本場無獨立二周目條件"
            else:
                subtitle = concise(s["objective"], 85)
            draw_box(d, (x, y, x+bw, y+bh), f"{s['id']}\n{scene_title(s)}", subtitle, kind, marker=status == "需補強")
            centers.append((x, y, x+bw, y+bh))
        for idx in range(len(centers)-1):
            a, b = centers[idx], centers[idx+1]
            if b[1] == a[1]:
                draw_arrow(d, (a[2], (a[1]+a[3])//2), (b[0], (b[1]+b[3])//2), COLORS["ink"] if cond == "first_run" else COLORS["blue"], width=5)
            else:
                draw_arrow(d, ((a[0]+a[2])//2, a[3]), ((b[0]+b[2])//2, b[1]), COLORS["ink"] if cond == "first_run" else COLORS["blue"], width=5)
    map_legend(d, 95, 2110)
    path = MAP_DIR / f"ch{chapter}_trunk.png"
    img.save(path, quality=95)
    return path


def create_branch_maps(chapter, scenes, side_stories):
    relevant = [s for s in side_stories if any(p.get("chapter") == chapter for p in s.get("stagePlan", []))]
    chunks = [relevant[i:i+10] for i in range(0, len(relevant), 10)] or [[]]
    paths = []
    anchors = {s["id"]: s for s in scenes}
    for part, chunk in enumerate(chunks, 1):
        w, h = 3600, 2250
        img = Image.new("RGB", (w, h), COLORS["paper"])
        d = ImageDraw.Draw(img)
        suffix = f"（{part}/{len(chunks)}）" if len(chunks) > 1 else ""
        map_header(d, f"第 {chapter} 章分支｜{CHAPTER_TITLES[chapter]}{suffix}", "人物支線、多人支線、Boss 線與回歸主幹的位置；待製作內容以灰線表示", w)
        d.text((1390, 245), "主線錨點", font=font(32), fill=COLORS["ink"])
        anchor_scenes = [scenes[0], scenes[len(scenes)//2], scenes[-1]]
        anchor_boxes = []
        for idx, s in enumerate(anchor_scenes):
            y = 330 + idx*580
            box = (1310, y, 2290, y+240)
            draw_box(d, box, f"{s['id']}｜{scene_title(s)}", concise(s["outputsRaw"], 120), "main")
            anchor_boxes.append((s, box))
            if idx:
                draw_arrow(d, (1800, anchor_boxes[idx-1][1][3]), (1800, box[1]), COLORS["ink"], width=7)
        for idx, story in enumerate(chunk):
            side = idx % 2
            row = idx // 2
            x1 = 90 if side == 0 else 2460
            x2 = 1140 if side == 0 else 3510
            y1 = 300 + row*350
            box = (x1, y1, x2, y1+265)
            stage = next(p for p in story.get("stagePlan", []) if p.get("chapter") == chapter)
            # Story category stays gold; the dotted connector and status copy
            # communicate that the approved stage plan is not implemented yet.
            kind = "side"
            title = f"{story['id']}｜{story['title']}"
            subtitle = f"入口：{story.get('unlockAfterSceneId') or '待定'}\n階段：{concise(stage.get('objective'), 95)}"
            draw_box(d, box, title, subtitle, kind)
            anchor_idx = min(row, 2)
            anchor = anchor_boxes[anchor_idx][1]
            if side == 0:
                draw_arrow(d, (box[2], (box[1]+box[3])//2), (anchor[0], (anchor[1]+anchor[3])//2), COLORS["gray"], width=5, dotted=True)
            else:
                draw_arrow(d, (box[0], (box[1]+box[3])//2), (anchor[2], (anchor[1]+anchor[3])//2), COLORS["gray"], width=5, dotted=True)
        main_id, main_name = MAIN_BOSSES[chapter]
        draw_box(d, (1310, 1890, 2290, 2070), f"章節收束 Boss｜{main_name}", main_id, "boss")
        draw_arrow(d, (1800, anchor_boxes[-1][1][3]), (1800, 1890), COLORS["red"], width=8)
        ext_id, ext_name = EXTERNAL_BOSSES[chapter]
        draw_box(d, (2460, 1890, 3510, 2070), f"二周目外部 Boss 保留位｜{ext_name}", "路線、獎勵、戰鬥、旗標與美術均暫緩。", "pending")
        map_legend(d, 95, 2110)
        path = MAP_DIR / f"ch{chapter}_branches_{part}.png"
        img.save(path, quality=95)
        paths.append(path)
    return paths


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=100, bottom=90, end=100):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, name="Noto Serif TC", size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color.replace("#", ""))


def configure_styles(doc):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Noto Serif TC"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Serif TC")
    normal.font.size = Pt(9.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.16
    for name, size, color in (("Title", 17, "252525"), ("Heading 1", 14, "252525"), ("Heading 2", 12, "7A2626"), ("Heading 3", 10.5, "246A92")):
        style = styles[name]
        style.font.name = "Noto Serif TC"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Serif TC")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
    if "Audit Label" not in styles:
        style = styles.add_style("Audit Label", WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles["Audit Label"]
    style.font.name = "Noto Sans TC"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Sans TC")
    style.font.size = Pt(8)
    style.font.bold = True
    style.font.color.rgb = RGBColor(102, 102, 102)
    style.paragraph_format.space_after = Pt(1)
    if "List Bullet" not in styles:
        bullet = styles.add_style("List Bullet", WD_STYLE_TYPE.PARAGRAPH)
        bullet.base_style = styles["Normal"]
        bullet.paragraph_format.left_indent = Inches(0.22)
        bullet.paragraph_format.first_line_indent = Inches(-0.14)
    else:
        bullet = styles["List Bullet"]
    bullet.font.name = "Noto Serif TC"
    bullet._element.rPr.rFonts.set(qn("w:eastAsia"), "Noto Serif TC")
    bullet.font.size = Pt(9.5)


def clear_template_body(doc):
    body = doc._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def add_rule(doc, color="B8B8B8", size="8"):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(6)
    p_pr = p._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)
    p_pr.append(borders)


def add_label_value(doc, label, value, status=None):
    p = doc.add_paragraph()
    p.paragraph_format.keep_together = True
    r = p.add_run(label + "　")
    set_run_font(r, "Noto Sans TC", 8.5, True, "666666")
    r = p.add_run(value)
    set_run_font(r, "Noto Serif TC", 9.5, False, "252525")
    if status:
        r = p.add_run("　" + status)
        color = "D08A13" if status == "需補強" else "8A8A8A" if status == "待定" else "246A92"
        set_run_font(r, "Noto Sans TC", 8.5, True, color)
    return p


def add_callout(doc, title, body, kind="finding"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.35)
    cell = table.cell(0, 0)
    set_cell_margins(cell, 130, 150, 130, 150)
    fill = {"finding": "F4F4F4", "weak": "FFF2D8", "pending": "EFEFEF", "risk": "F5E4E4"}[kind]
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    set_run_font(r, "Noto Sans TC", 9.5, True, "7A2626" if kind == "risk" else "252525")
    p = cell.add_paragraph(body)
    p.paragraph_format.space_after = Pt(0)
    for r in p.runs:
        set_run_font(r, "Noto Serif TC", 9.2)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_toc_field(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = ' TOC \\o "1-3" \\h \\z \\u '
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "目錄將在 Word 開啟或匯出時更新。"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, placeholder, fld_end])


def add_landscape_images(doc, image_entries):
    """Add consecutive diagram pages in one landscape section, then return to portrait."""
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    section.top_margin = Inches(0.42)
    section.bottom_margin = Inches(0.48)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)
    section.footer.is_linked_to_previous = True
    for idx, (image_path, title) in enumerate(image_entries):
        if idx:
            doc.add_page_break()
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run()
        r.add_picture(str(image_path), width=Inches(9.95))
        pic = p._p.xpath('.//wp:docPr')
        if pic:
            pic[0].set("descr", title)
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    section.footer.is_linked_to_previous = True


def build_doc(data, maps):
    shutil.copy2(TEMPLATE, OUTPUT)
    doc = Document(OUTPUT)
    clear_template_body(doc)
    configure_styles(doc)
    base = doc.sections[0]
    base.top_margin = Inches(0.7)
    base.bottom_margin = Inches(0.72)
    base.left_margin = Inches(0.85)
    base.right_margin = Inches(0.85)
    props = doc.core_properties
    props.title = "SDS_Games 雙周目劇情因果備忘錄"
    props.subject = "七章主線、兩周目、Boss、支線與角色交疊之因果審查"
    props.author = "OpenAI Codex（依專案權威資料編製）"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("INTERNAL STORY CAUSAL REVIEW")
    set_run_font(r, "Noto Sans TC", 8, True, "777777")
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("SDS_Games 雙周目劇情因果備忘錄")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("劇情因果審查稿｜非權威設定文件")
    set_run_font(r, "Noto Sans TC", 9.5, True, "7A2626")
    add_rule(doc, "7A2626", "12")
    meta = [
        ("檢視對象", "七章 66 場主線、兩周目差異、章節 Boss、33 條人物故事與 6 條多人支線"),
        ("編製依據", "MAIN_STORY_BIBLE、CHAPTER_QUEST_FRAMEWORK、NARRATIVE_WRITING_GUIDE、角色檔案與 runtime registries"),
        ("版本", date.today().isoformat()),
        ("檢視主題", "前因、資訊邊界、代價、輸出、支線回歸與兩周目資訊隔離"),
    ]
    table = doc.add_table(rows=len(meta), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(1.1)
    table.columns[1].width = Inches(5.3)
    for row, (label, value) in zip(table.rows, meta):
        row.cells[0].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        row.cells[1].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
        for c in row.cells:
            set_cell_margins(c, 35, 50, 35, 50)
        r = row.cells[0].paragraphs[0].add_run(label)
        set_run_font(r, "Noto Sans TC", 8.5, True, "666666")
        r = row.cells[1].paragraphs[0].add_run(value)
        set_run_font(r, "Noto Serif TC", 9.3)
    add_rule(doc)
    add_callout(doc, "審查定位", "本稿把 runtime 與權威文件轉成可逐場核對的閱讀稿。它不自動覆蓋 MAIN_STORY_BIBLE 或其他權威 MD；候選修補橋段均屬非定案建議。心智圖是高解析度靜態圖，不能在 Word 內拖曳節點。", "finding")

    doc.add_heading("摘要判定", level=1)
    md_ids = extract_md_scene_ids()
    runtime_ids = data["sceneOrder"]
    missing_in_md = sorted(set(runtime_ids)-set(md_ids))
    extra_in_md = sorted(set(md_ids)-set(runtime_ids))
    scene_statuses = Counter(audit_scene(s)[0] for s in data["scenes"])
    findings = [
        f"runtime 主線共 {len(runtime_ids)} 場，章節分布為 " + "、".join(f"第 {ch} 章 {len(data['sceneOrderByChapter'][str(ch)])} 場" for ch in range(1,8)) + "。",
        f"可選故事共 {len(data['sideStories'])} 條，其中人物故事 {len(data['personalSideStoryIds'])} 條、多人故事 {len(data['ensembleSideStoryIds'])} 條；目前全部屬 stage plan／待正式製作狀態。",
        f"逐場欄位審查：已確立 {scene_statuses['已確立']} 場、需補強 {scene_statuses['需補強']} 場。此判定只檢查因果接口完整度，不代表文本與演出已驗收。",
        "七個二周目外部 Boss 僅保留核准身分與容量邊界；路線、掉落、獎勵、戰鬥、旗標與美術均維持暫緩。",
    ]
    for item in findings:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run("• " + item)
    if missing_in_md or extra_in_md:
        add_callout(doc, "來源差異｜需人工確認", f"runtime 未在 MAIN_STORY_BIBLE 正規式掃描中找到：{', '.join(missing_in_md) or '無'}；文件多出：{', '.join(extra_in_md) or '無'}。本稿不靜默選邊。", "weak")
    else:
        add_callout(doc, "來源一致性", "MAIN_STORY_BIBLE 與 runtime registry 均可識別相同的 66 個場景 ID；場景順序以 runtime StorySceneOrder 呈現。", "finding")
    untitled = [s["id"] for s in data["scenes"] if not s.get("title")]
    if untitled:
        add_callout(doc, "顯示標題差異｜需補強", f"runtime 有 {len(untitled)} 場未提供獨立 `title` 欄位。本稿以「未定標題＋scene id 尾段」顯示，不把英文 id 尾段翻成正式中文標題。正式命名仍須回到權威資料核准。", "weak")

    doc.add_heading("閱讀方式", level=1)
    doc.add_paragraph("每章先看雙周目主幹圖，再看分支圖，最後逐場核對前因、人物目的、玩家理解、危機作用、輸出與回歸點。灰色虛線與「待定」不代表已批准；琥珀標記表示因果接口需要補強。")
    doc.add_heading("目錄", level=1)
    add_toc_field(doc)
    add_landscape_images(doc, [(maps["overview"], "七章雙周目因果總覽")])

    scenes_by_ch = defaultdict(list)
    for s in data["scenes"]:
        scenes_by_ch[s["chapter"]].append(s)
    sides_by_unlock = defaultdict(list)
    for story in data["sideStories"]:
        sides_by_unlock[story.get("unlockAfterSceneId")].append(story)

    for ch in range(1, 8):
        scenes = scenes_by_ch[ch]
        doc.add_heading(f"第 {ch} 章｜{CHAPTER_TITLES[ch]}", level=1)
        doc.add_paragraph(f"本章共有 {len(scenes)} 場主線。固定收束為 {MAIN_BOSSES[ch][1]}（{MAIN_BOSSES[ch][0]}）；本章 stage plan 涵蓋 {sum(1 for s in data['sideStories'] if any(p.get('chapter') == ch for p in s.get('stagePlan', [])))} 條可選故事。")
        diagram_entries = [(maps[f"ch{ch}_trunk"], f"第 {ch} 章雙周目主幹")]
        diagram_entries.extend((path, f"第 {ch} 章分支圖 {idx}") for idx, path in enumerate(maps[f"ch{ch}_branches"], 1))
        add_landscape_images(doc, diagram_entries)
        doc.add_heading("逐場因果審查", level=2)
        for idx, scene in enumerate(scenes):
            status, reason = audit_scene(scene)
            doc.add_heading(f"{scene['id']}｜{scene_title(scene)}", level=3)
            add_label_value(doc, "審查狀態", reason, status)
            add_label_value(doc, "前置事件與觸發", concise(scene.get("entry")))
            actor_ids = scene_actor_ids(scene)
            actor_text = "、".join(ACTOR_NAMES.get(a, a) for a in actor_ids) or concise(scene.get("participantsRaw"))
            add_label_value(doc, "當場人物／目的", actor_text + "；" + concise(scene.get("objective"), 280))
            add_label_value(doc, "已知資訊邊界", concise(scene.get("knowledgeBoundary") or scene.get("inputsRaw"), 300))
            add_label_value(doc, "玩家取得的證據或理解", concise(scene.get("outputsRaw"), 340))
            is_boss_scene = MAIN_BOSSES[ch][0] in (scene["id"] + " " + scene.get("outputsRaw", "")) or any(rb[2] == scene["id"] for rb in ROUTE_BOSSES[ch])
            crisis = "Boss／不可逆危機節點。" if is_boss_scene else "承擔調查、關係、路線或世界狀態推進。"
            add_label_value(doc, "Boss／危機的故事作用", crisis + " " + concise(scene.get("objective"), 240))
            next_scene = scenes[idx+1]["id"] if idx + 1 < len(scenes) else (data["sceneOrderByChapter"].get(str(ch+1), ["章末"])[0] if ch < 7 else "結局分流")
            add_label_value(doc, "場景輸出／下一場承接", concise(scene.get("outputsRaw"), 260) + f"；下一承接：{next_scene}。")
            profile = scene_run_profile(scene)
            if profile["first"] or profile["second"]:
                add_label_value(doc, "第一周目", beat_summary(scene, "first_run"))
                add_label_value(doc, "第二周目", beat_summary(scene, "second_run"))
            else:
                add_label_value(doc, "兩周目差異", "本場只有共用節拍；若兩周目需要不同理解，差異必須由相鄰場景、既有成就記憶或明示條件承擔。")
            branches = sides_by_unlock.get(scene["id"], [])
            if branches:
                branch_text = "；".join(f"{s['id']}（{s['title']}，{s['status']}）" for s in branches)
            else:
                branch_text = "本場沒有直接登記的可選故事入口。"
            add_label_value(doc, "可分岔支線／回歸", branch_text)
            doc.add_paragraph().paragraph_format.space_after = Pt(1)

    doc.add_page_break()
    doc.add_heading("39 條可選故事總表", level=1)
    doc.add_paragraph("以下內容依 OptionalSideStoryRegistry 與 OptionalEnsembleStoryRegistry 展開。所有條目均為已核准方向、待正式製作；本文不把 stage plan 偽裝成完整場景。")
    for story in data["sideStories"]:
        story_kind = "多人支線" if story["id"] in data["ensembleSideStoryIds"] else "人物故事"
        doc.add_heading(f"{story['id']}｜{story['title']}", level=2)
        add_label_value(doc, "狀態／類型", f"{story.get('status')}；{story_kind}", "待定")
        add_label_value(doc, "入口", story.get("unlockAfterSceneId") or "未明定入口")
        chars = "、".join(ACTOR_NAMES.get(x, x) for x in story.get("characterIds", []))
        add_label_value(doc, "人物與作用", f"{chars or ACTOR_NAMES.get(story.get('primaryCharacterId'), story.get('primaryCharacterId', '未定'))}；{concise(story.get('purpose'))}")
        add_label_value(doc, "人物揭示", concise(story.get("characterReveal")))
        for i, stage in enumerate(story.get("stagePlan", []), 1):
            add_label_value(doc, f"階段 {i}", f"第 {stage.get('chapter')} 章／{stage.get('ownerId')}：{concise(stage.get('objective'), 320)}")
        last_ch = max((p.get("chapter", 1) for p in story.get("stagePlan", [])), default=1)
        chapter_last = data["sceneOrderByChapter"].get(str(last_ch), ["未定"])[-1]
        add_label_value(doc, "回歸主幹或終點", f"registry 未提供獨立 rejoin flag；審查上暫以最後階段所在章的章末主線 {chapter_last} 作為回歸檢查點，正式製作時需明定。", "需補強")
        add_label_value(doc, "主線邊界", concise(story.get("mainlineBoundary")))
        resources = story.get("resourceNeeds", {})
        resource_text = []
        for key, label in (("newItemIds", "物品"), ("newIconIds", "圖示"), ("newBackgroundIds", "背景"), ("newSystemHooks", "系統接口")):
            vals = resources.get(key) or []
            if vals:
                resource_text.append(f"{label}：{', '.join(vals)}")
        add_label_value(doc, "資源／資料影響", "；".join(resource_text) if resource_text else "未要求新角色、怪物、素材、地點或圖片；仍需正式對話與演出驗證。")

    doc.add_page_break()
    doc.add_heading("Boss 因果索引", level=1)
    for ch in range(1, 8):
        main_id, main_name = MAIN_BOSSES[ch]
        relevant = [s for s in scenes_by_ch[ch] if main_id in (s["id"] + " " + s.get("outputsRaw", "") + " " + s.get("objective", ""))]
        doc.add_heading(f"第 {ch} 章｜{main_name}（{main_id}）", level=2)
        cause = concise(" ／ ".join(s.get("inputsRaw", "") for s in relevant), 400) if relevant else "runtime 場景未以 Boss id 明示形成原因，需由章節前置場景反查。"
        consequence = concise(" ／ ".join(s.get("outputsRaw", "") for s in relevant), 450) if relevant else "待人工確認。"
        chars = []
        for s in relevant:
            chars.extend(scene_actor_ids(s))
        add_label_value(doc, "形成原因", cause, "已確立" if relevant else "需補強")
        add_label_value(doc, "相關人物", "、".join(dict.fromkeys(ACTOR_NAMES.get(x, x) for x in chars)) or "未在 Boss 場景明列")
        add_label_value(doc, "擊敗／避免戰鬥後果", consequence)
        if ROUTE_BOSSES[ch]:
            add_label_value(doc, "路線 Boss", "；".join(f"{name}（{bid}，場景 {sid}）" for bid, name, sid in ROUTE_BOSSES[ch]))
        else:
            add_label_value(doc, "路線 Boss", "本章框架未明定獨立路線 Boss；不可自行補一隻。", "待定")
        ext_id, ext_name = EXTERNAL_BOSSES[ch]
        add_label_value(doc, "二周目外部 Boss 保留位", f"{ext_name}（{ext_id}）。僅核准身分與容量邊界；路線、獎勵、戰鬥、旗標與美術暫緩。", "待定")

    doc.add_page_break()
    doc.add_heading("角色交疊索引", level=1)
    doc.add_paragraph("本索引追蹤核心角色跨章節的投入、資訊邊界、代價與回收。它只以角色實際出現在 scene beats 或 side-story characterIds 為準。")
    for actor_id in CORE_CHARACTER_IDS:
        name = ACTOR_NAMES[actor_id]
        scenes = [s for s in data["scenes"] if actor_id in scene_actor_ids(s)]
        chapters = sorted({s["chapter"] for s in scenes})
        side = [s for s in data["sideStories"] if actor_id in s.get("characterIds", []) or actor_id == s.get("primaryCharacterId")]
        doc.add_heading(f"{name}｜{actor_id}", level=2)
        add_label_value(doc, "跨章投入", "、".join(f"第 {ch} 章" for ch in chapters) or "runtime beat 未直接登場")
        add_label_value(doc, "主線節點", "、".join(s["id"] for s in scenes) or "無")
        add_label_value(doc, "支線節點", "、".join(f"{s['id']}（{s['title']}）" for s in side) or "無")
        contracts = data.get("mainlineCharacterContracts", {}).get(actor_id)
        if contracts:
            add_label_value(doc, "角色契約", concise(json.dumps(contracts, ensure_ascii=False), 500), "已確立")
        else:
            add_label_value(doc, "角色契約", "MainlineCharacterContracts 未列出；以 scene beats 與角色檔案交叉審查。", "需補強")
        outputs = concise(" ／ ".join(s.get("outputsRaw", "") for s in scenes[-3:]), 500) if scenes else "無"
        add_label_value(doc, "代價與回收觀察", outputs)

    doc.add_page_break()
    doc.add_heading("薄弱節點與非定案修補候選", level=1)
    add_callout(doc, "重要邊界", "以下候選只提供因果修補方向，不構成正式設定。任何候選若需要新角色、怪物、素材、地點或圖片，必須先經使用者批准並同步資料與驗證計畫。", "risk")
    weak_scenes = [(s, audit_scene(s)[1]) for s in data["scenes"] if audit_scene(s)[0] == "需補強"]
    if weak_scenes:
        for s, reason in weak_scenes:
            doc.add_heading(f"{s['id']}｜{scene_title(s)}", level=2)
            add_label_value(doc, "問題", reason, "需補強")
            add_label_value(doc, "候選 A（非定案）", "把缺少的前因改成前一場已取得的證據或人物動機，不新增實體資源。")
            add_label_value(doc, "候選 B（非定案）", "把缺少的輸出明定為既有旗標、人物理解或路線狀態，避免另建第二套任務資料。")
            add_label_value(doc, "資源成本", "預設不新增角色、怪物、素材、地點或圖片；若現有場景無法承擔，須另提資源申請。")
    else:
        doc.add_paragraph("66 場均具備入口、目的、輸出與離場欄位；較大的未定風險集中在支線 rejoin 契約及二周目外部 Boss 暫緩層。")
    doc.add_heading("支線回歸接口", level=2)
    add_label_value(doc, "問題", "39 條支線都有入口與 stage plan，但 registry 沒有統一的顯式 rejoinSceneId／terminalState 欄位。", "需補強")
    add_label_value(doc, "候選 A（非定案）", "為所有支線增加 `rejoinSceneId` 或 `terminalState`，只描述回歸接口，不改支線內容。")
    add_label_value(doc, "候選 B（非定案）", "若支線可跨章，增加 `stageReturnFlags`，讓每一階段清楚回到哪個既有主線狀態。")
    add_label_value(doc, "資源成本", "資料 schema 與 validator 需調整；不需新增角色、怪物、素材、地點或圖片。")
    doc.add_heading("二周目外部 Boss", level=2)
    add_label_value(doc, "問題", "七個身分已核准，但所有路線、獎勵、戰鬥、旗標與美術都被刻意暫緩。", "待定")
    add_label_value(doc, "候選方向（非定案）", "維持灰色保留位，直到第一周目完整驗收；屆時逐章提出故事理由、來源、地圖分支、資料影響、資產需求與驗證方案。")
    add_label_value(doc, "資源成本", "尚未估算；不得在本稿中預先生成或配置。")

    doc.add_page_break()
    doc.add_heading("結論", level=1)
    doc.add_paragraph("現有 runtime 已提供可審查的 66 場主幹與 39 條可選故事 stage plan。主線場景的資料接口大致完整，真正需要優先補上的不是更多劇情，而是支線回歸契約、來源差異的人工判讀，以及 Chapter 1–2 無跳過實機驗收。二周目外部 Boss 應繼續維持暫緩，避免尚未驗收的第一周目核心被外部擴充反向牽動。")
    doc.add_heading("驗證清單", level=2)
    checks = [
        f"66 個主線場景：{len(set(data['sceneOrder']))}/66，無重複。",
        f"39 條可選故事：{len({s['id'] for s in data['sideStories']})}/39；每條均列入口、階段與暫定回歸檢查點。",
        "七章主 Boss 與七個二周目外部 Boss 保留位：完整。",
        "正式設定、需補強、待定候選：已使用不同標記分離。",
        "兩周目資訊：依 beat condition 分列；共用節拍不偽造差異。",
        "候選若涉及新資源：已要求另列成本與資料影響。",
    ]
    for check in checks:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run("• " + check)
    doc.add_heading("資料來源", level=2)
    for source in [
        "docs/MAIN_STORY_BIBLE.md",
        "docs/CHAPTER_QUEST_FRAMEWORK.md",
        "docs/NARRATIVE_WRITING_GUIDE.md",
        "docs/characters/*.md",
        "src/js/data/StorySceneRegistry.js",
        "src/js/data/OptionalSideStoryRegistry.js",
        "src/js/data/StoryActors.js",
    ]:
        doc.add_paragraph("• " + source, style="List Bullet")

    settings = doc.settings._element
    # The retained template requests embedded TrueType fonts. With large CJK
    # variable fonts Word expands the DOCX dramatically and PDF export can take
    # many minutes. Keep the requested font names, but rely on the installed
    # Noto fonts instead of embedding them into this local review artifact.
    for tag in ("embedTrueTypeFonts", "embedSystemFonts", "saveSubsetFonts"):
        node = settings.find(qn(f"w:{tag}"))
        if node is not None:
            settings.remove(node)
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")
    doc.save(OUTPUT)


def validate(data, maps):
    assert len(data["scenes"]) == 66
    assert len(set(data["sceneOrder"])) == 66
    assert len(data["sideStories"]) == 39
    assert len(set(s["id"] for s in data["sideStories"])) == 39
    assert len(MAIN_BOSSES) == 7 and len(EXTERNAL_BOSSES) == 7
    assert len([maps["overview"]] + [maps[f"ch{c}_trunk"] for c in range(1,8)] + [p for c in range(1,8) for p in maps[f"ch{c}_branches"]]) >= 15
    for s in data["scenes"]:
        assert s.get("entry") and s.get("outputsRaw") and s.get("exit"), s["id"]
    for story in data["sideStories"]:
        assert story.get("stagePlan"), story["id"]
        assert story.get("unlockAfterSceneId") or story.get("offerPolicy"), story["id"]


def main():
    TMP.mkdir(parents=True, exist_ok=True)
    MAP_DIR.mkdir(parents=True, exist_ok=True)
    data = load_data()
    scenes_by_ch = defaultdict(list)
    for scene in data["scenes"]:
        scenes_by_ch[scene["chapter"]].append(scene)
    maps = {"overview": create_overview_map(data)}
    for ch in range(1, 8):
        maps[f"ch{ch}_trunk"] = create_trunk_map(ch, scenes_by_ch[ch])
        maps[f"ch{ch}_branches"] = create_branch_maps(ch, scenes_by_ch[ch], data["sideStories"])
    validate(data, maps)
    build_doc(data, maps)
    print(json.dumps({
        "output": str(OUTPUT),
        "scenes": len(data["scenes"]),
        "sideStories": len(data["sideStories"]),
        "maps": 1 + 7 + sum(len(maps[f"ch{c}_branches"]) for c in range(1,8)),
        "size": OUTPUT.stat().st_size,
    }, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
